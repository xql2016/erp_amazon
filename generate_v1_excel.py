#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
V1功能Excel文件生成脚本
用于生成控制流量和导入流量的MSKU清单Excel模板

使用方法:
    pip install openpyxl
    python generate_v1_excel.py
"""

import openpyxl
from openpyxl import Workbook

def generate_control_traffic_excel():
    """生成控制流量MSKU清单.xlsx"""
    wb = Workbook()
    ws = wb.active
    ws.title = "控制流量MSKU"
    
    # 第一行：站点ProfileId（示例数据，请替换为实际ProfileId）
    profile_ids = ["123456789", "234567890", "345678901", "456789012", "567890123"]
    for col, profile_id in enumerate(profile_ids, start=1):
        ws.cell(row=1, column=col, value=profile_id)
    
    # 第二行开始：MSKU清单（示例数据，请替换为实际MSKU）
    msku_data = [
        ["CTRL-DE-001", "CTRL-FR-001", "CTRL-IT-001", "CTRL-ES-001", "CTRL-UK-001"],
        ["CTRL-DE-002", "CTRL-FR-002", "CTRL-IT-002", "CTRL-ES-002", "CTRL-UK-002"],
        ["CTRL-DE-003", "CTRL-FR-003", "CTRL-IT-003", "CTRL-ES-003", "CTRL-UK-003"],
        ["CTRL-DE-004", "CTRL-FR-004", "CTRL-IT-004", "CTRL-ES-004", "CTRL-UK-004"],
        ["CTRL-DE-005", "CTRL-FR-005", "CTRL-IT-005", "CTRL-ES-005", "CTRL-UK-005"],
        ["CTRL-DE-006", "CTRL-FR-006", "CTRL-IT-006", "CTRL-ES-006", "CTRL-UK-006"],
        ["CTRL-DE-007", "CTRL-FR-007", "CTRL-IT-007", "CTRL-ES-007", "CTRL-UK-007"],
        ["CTRL-DE-008", "CTRL-FR-008", "CTRL-IT-008", "CTRL-ES-008", "CTRL-UK-008"],
        ["CTRL-DE-009", "CTRL-FR-009", "CTRL-IT-009", "CTRL-ES-009", "CTRL-UK-009"],
        ["CTRL-DE-010", "CTRL-FR-010", "CTRL-IT-010", "CTRL-ES-010", "CTRL-UK-010"],
    ]
    
    for row_idx, row_data in enumerate(msku_data, start=2):
        for col_idx, msku in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=msku)
    
    # 调整列宽
    for col in range(1, len(profile_ids) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    
    # 保存文件
    file_path = "src/main/resources/控制流量MSKU清单.xlsx"
    wb.save(file_path)
    print(f"✓ {file_path} 生成成功！")
    return file_path

def generate_import_traffic_excel():
    """生成导入流量MSKU清单.xlsx"""
    wb = Workbook()
    ws = wb.active
    ws.title = "导入流量MSKU"
    
    # 第一行：站点ProfileId（示例数据，请替换为实际ProfileId）
    profile_ids = ["123456789", "234567890", "345678901", "456789012", "567890123"]
    for col, profile_id in enumerate(profile_ids, start=1):
        ws.cell(row=1, column=col, value=profile_id)
    
    # 第二行开始：MSKU清单（示例数据，请替换为实际MSKU）
    msku_data = [
        ["IMPT-DE-001", "IMPT-FR-001", "IMPT-IT-001", "IMPT-ES-001", "IMPT-UK-001"],
        ["IMPT-DE-002", "IMPT-FR-002", "IMPT-IT-002", "IMPT-ES-002", "IMPT-UK-002"],
        ["IMPT-DE-003", "IMPT-FR-003", "IMPT-IT-003", "IMPT-ES-003", "IMPT-UK-003"],
        ["IMPT-DE-004", "IMPT-FR-004", "IMPT-IT-004", "IMPT-ES-004", "IMPT-UK-004"],
        ["IMPT-DE-005", "IMPT-FR-005", "IMPT-IT-005", "IMPT-ES-005", "IMPT-UK-005"],
        ["IMPT-DE-006", "IMPT-FR-006", "IMPT-IT-006", "IMPT-ES-006", "IMPT-UK-006"],
        ["IMPT-DE-007", "IMPT-FR-007", "IMPT-IT-007", "IMPT-ES-007", "IMPT-UK-007"],
        ["IMPT-DE-008", "IMPT-FR-008", "IMPT-IT-008", "IMPT-ES-008", "IMPT-UK-008"],
        ["IMPT-DE-009", "IMPT-FR-009", "IMPT-IT-009", "IMPT-ES-009", "IMPT-UK-009"],
        ["IMPT-DE-010", "IMPT-FR-010", "IMPT-IT-010", "IMPT-ES-010", "IMPT-UK-010"],
        ["IMPT-DE-011", "IMPT-FR-011", "IMPT-IT-011", "IMPT-ES-011", "IMPT-UK-011"],
        ["IMPT-DE-012", "IMPT-FR-012", "IMPT-IT-012", "IMPT-ES-012", "IMPT-UK-012"],
    ]
    
    for row_idx, row_data in enumerate(msku_data, start=2):
        for col_idx, msku in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=msku)
    
    # 调整列宽
    for col in range(1, len(profile_ids) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    
    # 保存文件
    file_path = "src/main/resources/导入流量MSKU清单.xlsx"
    wb.save(file_path)
    print(f"✓ {file_path} 生成成功！")
    return file_path

def main():
    """主函数"""
    print("=" * 60)
    print("V1功能Excel文件生成工具")
    print("=" * 60)
    print("\n开始生成Excel文件...\n")
    
    try:
        # 生成控制流量Excel
        control_file = generate_control_traffic_excel()
        
        # 生成导入流量Excel
        import_file = generate_import_traffic_excel()
        
        print("\n" + "=" * 60)
        print("✓ 所有Excel文件生成成功！")
        print("=" * 60)
        print("\n生成的文件:")
        print(f"  1. {control_file}")
        print(f"  2. {import_file}")
        print("\n使用说明:")
        print("  1. 请将Excel文件中的ProfileId替换为实际站点ID")
        print("  2. 请将MSKU替换为实际需要处理的SKU")
        print("  3. 确保文件已放置在 src/main/resources/ 目录下")
        print("  4. 在配置文件中指定正确的文件路径")
        print("\nExcel文件格式:")
        print("  - 第一行: 站点ProfileId（数字）")
        print("  - 第二行开始: 每列对应站点的MSKU清单")
        print("  - 空单元格会被自动跳过")
        print("\n" + "=" * 60)
        
    except Exception as e:
        print(f"\n✗ 生成失败: {e}")
        print("\n请确保:")
        print("  1. 已安装openpyxl: pip install openpyxl")
        print("  2. src/main/resources/ 目录存在")
        print("  3. 有文件写入权限")

if __name__ == "__main__":
    main()

